"""报告生成层：Excel 对比报告 + Markdown 差异报告。

Excel 结构（对齐项目工程师的评审习惯）：
  Sheet1 汇总总览 —— 每份规格书的差异统计与风险分布
  Sheet2 差异清单 —— 所有 非一致 参数，红色=不满足 黄色=需评估
  Sheet3 参数明细 —— 全量参数逐条对比（含证据页码、置信度、双通道来源）
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ComparisonResult

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
RED = PatternFill("solid", fgColor="F8CBAD")
YELLOW = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6EFCE")

STATUS_FILL = {"不满足": RED, "需评估": YELLOW, "满足": GREEN}


def save_excel(results: list[ComparisonResult], path: str | Path) -> Path:
    path = Path(path)
    wb = Workbook()

    # ---- Sheet1 汇总 ----
    ws = wb.active
    ws.title = "汇总总览"
    ws.append(["规格书", "客户", "产品", "参数总数", "满足", "需评估", "不满足", "缺失", "高于(严于基线)", "低于(宽松)", "不同", "待人工复核"])
    by_spec: dict[str, list[ComparisonResult]] = defaultdict(list)
    for r in results:
        by_spec[r.spec_id].append(r)
    for spec_id in sorted(by_spec):
        rs = by_spec[spec_id]
        n = len(rs)
        ws.append(
            [
                spec_id,
                rs[0].customer_name,
                rs[0].product,
                n,
                sum(1 for r in rs if r.status == "满足"),
                sum(1 for r in rs if r.status == "需评估"),
                sum(1 for r in rs if r.status == "不满足"),
                sum(1 for r in rs if r.status == "缺失"),
                sum(1 for r in rs if r.direction == "高于"),
                sum(1 for r in rs if r.direction == "低于"),
                sum(1 for r in rs if r.direction == "不同"),
                sum(1 for r in rs if r.needs_review),
            ]
        )
    _style_header(ws)

    # ---- Sheet2 差异清单 ----
    ws2 = wb.create_sheet("差异清单")
    ws2.append(
        ["规格书", "客户", "参数编号", "参数名称", "类别", "客户要求", "内部基线", "单位", "差异方向", "判定", "差异说明", "证据页"]
    )
    for r in sorted(results, key=lambda x: (x.spec_id, x.param_id)):
        if r.direction == "一致":
            continue
        ws2.append(
            [r.spec_id, r.customer_name, r.param_id, r.param_name, r.category,
             r.customer_value, r.baseline_value, r.unit, r.direction, r.status,
             r.explanation, r.page]
        )
        cell = ws2.cell(row=ws2.max_row, column=10)
        if r.status in STATUS_FILL:
            cell.fill = STATUS_FILL[r.status]
    _style_header(ws2)

    # ---- Sheet3 参数明细 ----
    ws3 = wb.create_sheet("参数明细")
    ws3.append(
        ["规格书", "客户", "参数编号", "参数名称", "类别", "客户要求原文", "客户要求(归一)", "内部基线",
         "单位", "差异方向", "判定", "差异说明", "证据页", "置信度", "来源", "待复核"]
    )
    for r in sorted(results, key=lambda x: (x.spec_id, x.param_id)):
        ws3.append(
            [r.spec_id, r.customer_name, r.param_id, r.param_name, r.category, r.customer_raw,
             r.customer_value, r.baseline_value, r.unit, r.direction, r.status, r.explanation,
             r.page, r.confidence, r.source, "是" if r.needs_review else ""]
        )
        cell = ws3.cell(row=ws3.max_row, column=11)
        if r.status in STATUS_FILL:
            cell.fill = STATUS_FILL[r.status]
    _style_header(ws3)

    for sheet in (ws, ws2, ws3):
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = max(
                10, min(46, sheet.column_dimensions[get_column_letter(col)].width or 0)
            )
        sheet.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def save_markdown(results: list[ComparisonResult], path: str | Path) -> Path:
    path = Path(path)
    by_spec: dict[str, list[ComparisonResult]] = defaultdict(list)
    for r in results:
        by_spec[r.spec_id].append(r)

    lines = ["# 客户技术规格书差异分析报告", ""]
    for spec_id in sorted(by_spec):
        rs = by_spec[spec_id]
        gaps = [r for r in rs if r.direction != "一致"]
        hard = [r for r in gaps if r.status == "不满足"]
        review = [r for r in gaps if r.status == "需评估"]
        lines += [
            f"## {spec_id} · {rs[0].customer_name}（{rs[0].product}）",
            "",
            f"- 参数总数 **{len(rs)}**，其中差异 **{len(gaps)}** 项："
            f"不满足 **{len(hard)}**（能力缺口）、需评估 **{len(review)}**、宽松 **{len(gaps) - len(hard) - len(review)}**（降本空间）",
            "",
            "| 参数 | 客户要求 | 内部基线 | 方向 | 判定 | 说明 |",
            "| --- | --- | --- | --- | --- | --- |",
        ]
        for r in sorted(gaps, key=lambda x: (x.status, x.param_id)):
            lines.append(
                f"| {r.param_id} {r.param_name} | {r.customer_value} {r.unit} | {r.baseline_value} {r.unit} "
                f"| {r.direction} | **{r.status}** | {r.explanation} |"
            )
        lines.append("")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
